"""Grounded chat and access boundaries; synthetic policies are never real guidance."""

import hashlib
import json

from docx import Document
import pytest
from reportlab.pdfgen import canvas

from backend.auth.models import UserProfile
from backend.auth.policy_access import policy_scope
from rag_pipeline.documents.build import prepare_source_index
from rag_pipeline.documents.models import AccessScope
from rag_pipeline.documents.store import PolicyIndex, write_source_index
from rag_pipeline.retriever.section_retriever import SectionRetriever


@pytest.fixture
def corpus(tmp_path):
    source = tmp_path / "sources"
    source.mkdir()
    entries = []
    specs = [
        (
            "public-policy",
            "public",
            "active",
            "I. PURCHASING",
            "The departmental requester must document price reasonableness using comparable supplier quotations before approving the purchase. Refer to SUPPORT FORM for the supporting details.",
        ),
        (
            "support-form",
            "public",
            "active",
            "SUPPORT FORM",
            "Record the product description, supplier name and justification on this form. Include supporting quotations with the submission. This form records purchase documentation for auditing.",
        ),
        (
            "internal-policy",
            "internal",
            "active",
            "I. INTERNAL PRICING",
            "INTERNAL SECRET ONLY: the special internal pricing review requires confidential departmental approval before purchasing the restricted item from any approved vendor.",
        ),
        (
            "restricted-policy",
            "restricted",
            "active",
            "I. RESTRICTED PRICING",
            "RESTRICTED SECRET ONLY: the restricted supplier pricing agreement requires additional authorization before placing any purchase with this confidential contracted vendor.",
        ),
        (
            "draft-policy",
            "public",
            "review",
            "I. DRAFT",
            "DRAFT SECRET ONLY: this unreviewed supplier pricing guidance has not been released and must not be used for current procurement determinations.",
        ),
    ]
    for identifier, classification, status, heading, text in specs:
        path = source / f"{identifier}.docx"
        doc = Document()
        doc.add_heading(heading, 1)
        doc.add_paragraph(text)
        doc.save(path)
        entries.append(
            dict(
                id=identifier,
                filenames=[path.name],
                title=identifier,
                aliases=["SUPPORT FORM"] if identifier == "support-form" else [],
                kind="form" if identifier == "support-form" else "policy",
                access_classification=classification,
                review_status=status,
                reviewed_by="fixture",
                reviewed_at="2026-01-01T00:00:00Z",
                reviewed_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
            )
        )
    pdf = source / "pdf-policy.pdf"
    c = canvas.Canvas(str(pdf))
    c.drawString(30, 780, "I. COMPETITIVE PROPOSALS")
    c.drawString(
        30,
        750,
        "For the synthetic fixture purchase, obtain quotations from three qualified suppliers.",
    )
    c.save()
    entries.append(
        dict(
            id="pdf-policy",
            filenames=[pdf.name],
            title="PDF fixture",
            access_classification="public",
            review_status="active",
            reviewed_by="fixture",
            reviewed_at="2026-01-01T00:00:00Z",
            reviewed_sha256=hashlib.sha256(pdf.read_bytes()).hexdigest(),
        )
    )
    manifest = tmp_path / "manifest.json"
    manifest.write_text(json.dumps(dict(schema_version=1, documents=entries)))
    bundle = prepare_source_index(source, manifest, max_tokens=40, overlap_tokens=5)
    db = tmp_path / "sources.sqlite"
    write_source_index(bundle, db)
    return db, bundle, manifest, source


@pytest.mark.parametrize(
    "clearance,expected",
    [
        ("none", ("public",)),
        ("public", ("public",)),
        ("internal", ("public", "internal")),
        ("restricted", ("public", "internal", "restricted")),
        ("admin", ("public",)),
        ("department_buyer", ("public",)),
        ("INTERNAL", ("public",)),
    ],
)
def test_exact_server_clearance_mapping(clearance, expected):
    user = UserProfile(
        uid="u",
        email="u@ucsc.edu",
        is_approved=True,
        clearance=clearance,
        role="admin",
        department="IT",
    )
    scope = policy_scope(user)
    assert scope.classifications == expected
    assert not scope.include_review and not scope.include_archived


def test_pending_profile_does_not_grant_retrieval():
    with pytest.raises(PermissionError):
        policy_scope(UserProfile(uid="u", email="u@ucsc.edu", clearance="restricted"))


@pytest.mark.parametrize(
    "classes,allowed",
    [
        (("public",), {"public-policy", "support-form", "pdf-policy"}),
        (
            ("public", "internal"),
            {"public-policy", "support-form", "pdf-policy", "internal-policy"},
        ),
        (
            ("public", "internal", "restricted"),
            {
                "public-policy",
                "support-form",
                "pdf-policy",
                "internal-policy",
                "restricted-policy",
            },
        ),
    ],
)
def test_scope_filters_before_all_model_hooks_and_graph_expansion(
    corpus, classes, allowed
):
    db, _, _, _ = corpus
    seen = []

    def semantic(question, records):
        seen.extend(text for _, text in records)
        return [0.8] * len(records)

    def rerank(question, texts):
        seen.extend(texts)
        return list(range(len(texts), 0, -1))

    retriever = SectionRetriever(
        PolicyIndex(db, AccessScope(classifications=classes)), semantic, rerank
    )
    passages = retriever.retrieve("pricing purchase quotations")
    assert {doc.id for doc in retriever.corpus["documents"]} == allowed
    assert passages
    assert all(p.document.id in allowed for p in passages)
    text = " ".join(seen)
    assert "DRAFT SECRET ONLY" not in text
    if "internal" not in classes:
        assert "INTERNAL SECRET ONLY" not in text
    if "restricted" not in classes:
        assert "RESTRICTED SECRET ONLY" not in text
    allowed_versions = {doc.version_id for doc in retriever.corpus["documents"]}
    assert all(
        e.document_version_id in allowed_versions
        for edge in retriever.corpus["edges"]
        for e in edge.evidence
    )


def test_heading_match_returns_substantive_full_parent(corpus):
    db, _, _, _ = corpus
    hits = SectionRetriever(PolicyIndex(db)).retrieve("price reasonableness")
    assert hits
    assert "comparable supplier quotations" in hits[0].text
    assert len(hits[0].text) > 100
    assert hits[0].start == 0 and hits[0].end == len(hits[0].section.text)


def test_unknown_query_has_no_lexical_results_and_context_is_bounded(corpus):
    db, _, _, _ = corpus
    r = SectionRetriever(PolicyIndex(db))
    assert r.retrieve("astronauts nebula galaxies") == []
    hits = r.retrieve("quotations purchase", max_sources=2, max_characters=1000)
    assert len(hits) <= 2 and sum(len(p.text) for p in hits) <= 1000
    assert all(p.text == p.section.text[p.start : p.end] for p in hits)


def test_archive_hash_pins_current_version_and_blocks_unreviewed_replacement(corpus):
    _, _, manifest, source = corpus
    data = json.loads(manifest.read_text())
    e = next(e for e in data["documents"] if e["id"] == "public-policy")
    old = source / e["filenames"][0]
    e["archived_sha256"] = [e["reviewed_sha256"]]
    current = source / "public-current.docx"
    doc = Document()
    doc.add_paragraph("Current verified content only.")
    doc.save(current)
    e["filenames"].append(current.name)
    e["reviewed_sha256"] = hashlib.sha256(current.read_bytes()).hexdigest()
    manifest.write_text(json.dumps(data))
    bundle = prepare_source_index(source, manifest)
    old_doc = next(d for d in bundle.documents if d.file == old.name)
    assert old_doc.review_status == "archived"
    assert (
        next(d for d in bundle.documents if d.file == current.name).review_status
        == "active"
    )
    current.write_bytes(current.read_bytes() + b"unreviewed")
    with pytest.raises(ValueError, match="changed after review"):
        prepare_source_index(source, manifest)


def test_explicit_form_reference_expands_once_with_exact_source_evidence(corpus):
    db, _, _, _ = corpus
    r = SectionRetriever(PolicyIndex(db))
    hits = r.retrieve("departmental requester price reasonableness")
    linked = [p for p in hits if p.document.id == "support-form" and p.via_reference]
    assert linked
    ref = linked[0].via_reference
    assert ref["relation"] == "references" and ref["quote"] == "SUPPORT FORM"
    original = r.sections[ref["from_section_id"]]
    assert ref["quote"] in original.text
    assert ref["locations"]
    assert len(hits) <= 6
