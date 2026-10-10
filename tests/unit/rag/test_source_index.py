import hashlib
import json

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pytest
from pydantic import ValidationError
from reportlab.pdfgen import canvas

from rag_pipeline.documents.build import prepare_source_index
from rag_pipeline.documents.manifest import DocumentManifest
from rag_pipeline.documents.models import LOCAL_REVIEW_SCOPE, SourceLocation
from rag_pipeline.documents.store import PolicyIndex, write_source_index
from rag_pipeline.graph.schema import PolicyEdge
from rag_pipeline.graph.traversal import traverse
from rag_pipeline.indexer.source_chroma import (
    active_source_index_path,
    chroma_records,
    embedding_chunks,
    publish_chroma,
)


def manifest_for(tmp_path, entries):
    path = tmp_path / "manifest.json"
    path.write_text(json.dumps({"schema_version": 1, "documents": entries}))
    return path


def entry(path, identifier, classification="public", status="active", **extra):
    return {
        "id": identifier,
        "filenames": [path.name],
        "title": identifier,
        "policy_id": identifier.upper(),
        "aliases": [identifier.upper()],
        "kind": "policy",
        "access_classification": classification,
        "review_status": status,
        "reviewed_by": "fixture-reviewer",
        "reviewed_at": "2026-01-01T00:00:00Z",
        "reviewed_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        **extra,
    }


@pytest.fixture
def sources(tmp_path):
    source = tmp_path / "sources"
    source.mkdir()
    alpha = source / "Policy-Alpha.docx"
    doc = Document()
    doc.add_heading("I. GENERAL", level=1)
    doc.add_paragraph("Effective Date: 1/10/2026")
    doc.add_paragraph("See POL-B and FAR 2.101. A citation is not an override.")
    doc.add_heading("A. EXCEPTIONS", level=2)
    doc.add_paragraph("SHORT CONDITION MUST SURVIVE.")
    doc.add_paragraph(" ".join(f"word{i}" for i in range(1200)) + " FINAL-TAIL-MARKER")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Condition"
    table.cell(0, 1).text = "Required form"
    table.cell(1, 0).text = "Named exception"
    table.cell(1, 1).text = "SPECIAL-FORM-FIELD"
    p = doc.add_paragraph("Paragraph with hyperlink: ")
    hyperlink = OxmlElement("w:hyperlink")
    relationship = doc.part.relate_to(
        "https://example.invalid/reference",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "HYPERLINK-LABEL"
    run.append(text)
    hyperlink.append(run)
    p._p.append(hyperlink)
    box = OxmlElement("w:txbxContent")
    box_p = OxmlElement("w:p")
    box_r = OxmlElement("w:r")
    box_t = OxmlElement("w:t")
    box_t.text = "TEXTBOX-ONLY-REQUIREMENT"
    box_r.append(box_t)
    box_p.append(box_r)
    box.append(box_p)
    doc.add_paragraph("Textbox anchor")._p.append(box)
    doc.sections[0].header.paragraphs[0].text = "SOURCE HEADER"
    doc.sections[0].footer.paragraphs[0].text = "Revised 12-2025"
    doc.save(alpha)
    beta = source / "Policy-Beta.docx"
    doc = Document()
    doc.add_heading("I. INTERNAL", level=1)
    doc.add_paragraph("SECRET-BETA-ONLY clause.")
    doc.save(beta)
    form = source / "Form.docx"
    doc = Document()
    doc.add_heading("SPECIAL FORM", level=1)
    doc.add_paragraph("Form field and instructions")
    doc.save(form)
    manifest = manifest_for(
        tmp_path,
        [
            entry(alpha, "pol-a"),
            entry(beta, "pol-b", "internal"),
            entry(form, "special-form", kind="form", aliases=["SPECIAL FORM"]),
        ],
    )
    return source, manifest


def test_docx_preserves_tables_boxes_hyperlinks_dates_and_real_locations(sources):
    bundle = prepare_source_index(*sources, max_tokens=60, overlap_tokens=10)
    document = next(doc for doc in bundle.documents if doc.id == "pol-a")
    text = "\n".join(block.text for block in document.blocks)
    for marker in [
        "SHORT CONDITION MUST SURVIVE",
        "FINAL-TAIL-MARKER",
        "SPECIAL-FORM-FIELD",
        "HYPERLINK-LABEL",
        "TEXTBOX-ONLY-REQUIREMENT",
        "SOURCE HEADER",
        "Revised 12-2025",
    ]:
        assert marker in text
        assert marker in "\n".join(section.text for section in bundle.sections)
    assert any(block.kind == "textbox" for block in document.blocks)
    assert any(
        link.target == "https://example.invalid/reference"
        for block in document.blocks
        for link in block.links
    )
    table = next(block for block in document.blocks if block.kind == "table")
    assert table.rows[1][1].text == "SPECIAL-FORM-FIELD"
    assert all(block.location.page is None for block in document.blocks)
    assert any(
        date.field == "effective_date" and date.value == "2026-01-10"
        for date in document.dates
    )
    assert any(
        date.field == "revision_date"
        and date.value == "2025-12"
        and date.location.part.startswith("word/footer")
        for date in document.dates
    )


def test_full_parents_include_child_sections_and_every_chunk_has_exact_evidence(
    sources,
):
    bundle = prepare_source_index(*sources, max_tokens=60, overlap_tokens=10)
    sections = {section.id: section for section in bundle.sections}
    assert any(section.parent_section_id for section in bundle.sections)
    general = next(
        section for section in bundle.sections if section.title == "I. GENERAL"
    )
    assert "A. EXCEPTIONS" in general.text and "FINAL-TAIL-MARKER" in general.text
    assert "SHORT CONDITION MUST SURVIVE" in "\n".join(
        chunk.text for chunk in bundle.chunks
    )
    assert "FINAL-TAIL-MARKER" in "\n".join(chunk.text for chunk in bundle.chunks)
    for chunk in bundle.chunks:
        assert (
            chunk.text
            == sections[chunk.parent_section_id].text[chunk.start : chunk.end]
        )
        assert chunk.token_count <= 60
        assert chunk.locations
    for edge in bundle.edges:
        for evidence in edge.evidence:
            assert (
                evidence.quote
                == sections[evidence.section_id].text[evidence.start : evidence.end]
            )
    bundle.validate_integrity()


def test_ids_repeat_for_same_bytes_and_change_with_a_source_revision(sources):
    a, b = prepare_source_index(*sources), prepare_source_index(*sources)
    assert [doc.version_id for doc in a.documents] == [
        doc.version_id for doc in b.documents
    ]
    assert [chunk.id for chunk in a.chunks] == [chunk.id for chunk in b.chunks]
    assert [edge.id for edge in a.edges] == [edge.id for edge in b.edges]
    file = sources[0] / "Policy-Alpha.docx"
    doc = Document(file)
    doc.add_paragraph("Changed source")
    doc.save(file)
    with pytest.raises(ValueError, match="changed after review"):
        prepare_source_index(*sources)
    manifest = json.loads(sources[1].read_text())
    manifest["documents"][0]["review_status"] = "review"
    sources[1].write_text(json.dumps(manifest))
    changed = prepare_source_index(*sources)
    before = next(document for document in a.documents if document.id == "pol-a")
    after = next(document for document in changed.documents if document.id == "pol-a")
    assert before.id == after.id
    assert before.version_id != after.version_id


def test_filename_punctuation_changes_do_not_break_document_identity(sources):
    file = sources[0] / "Policy-Alpha.docx"
    file.rename(sources[0] / "Policy Alpha.docx")
    bundle = prepare_source_index(*sources)
    assert any(document.id == "pol-a" for document in bundle.documents)


def test_references_forms_and_external_stubs_have_evidence_without_inferred_overrides(
    sources,
):
    bundle = prepare_source_index(*sources)
    node_map = {node.id: node for node in bundle.nodes}
    references = [edge for edge in bundle.edges if edge.relation == "references"]
    assert any(
        node_map[edge.target_id].properties.get("document_id") == "pol-b"
        for edge in references
    )
    assert any(
        node_map[edge.target_id].kind == "reference"
        and edge.evidence[0].quote == "FAR 2.101"
        for edge in references
    )
    assert any(node.kind == "form" for node in bundle.nodes)
    assert any(edge.relation == "has_source" for edge in bundle.edges)
    assert not any(edge.relation == "overrides" for edge in bundle.edges)
    assert any(
        edge.evidence[0].link_target == "https://example.invalid/reference"
        and edge.evidence[0].quote == "HYPERLINK-LABEL"
        for edge in references
    )


def test_override_requires_conditions_review_and_valid_source_evidence(sources):
    bundle = prepare_source_index(*sources)
    reference = next(edge for edge in bundle.edges if edge.relation == "references")
    record = reference.model_dump()
    record["relation"] = "overrides"
    with pytest.raises(ValidationError, match="Overrides require"):
        PolicyEdge.model_validate(record)
    record.update(
        conditions={"funding": "fixture-only"},
        verified_by="reviewer",
        verified_at="2026-01-01T00:00:00Z",
    )
    verified = PolicyEdge.model_validate(record)
    assert verified.conditions == {"funding": "fixture-only"}
    record["evidence"] = []
    with pytest.raises(ValidationError, match="supporting source"):
        PolicyEdge.model_validate(record)


def test_store_search_parents_and_traversal_filter_internal_data_before_return(
    sources, tmp_path
):
    bundle = prepare_source_index(*sources)
    path = tmp_path / "index.sqlite"
    write_source_index(bundle, path)
    public = PolicyIndex(path)
    internal = PolicyIndex(path, LOCAL_REVIEW_SCOPE)
    assert {doc.id for doc in public.documents()} == {"pol-a", "special-form"}
    assert public.search("SECRET BETA") == []
    secret = internal.search("SECRET BETA")[0]
    assert public.chunk(secret.id) is None
    assert public.parent_section(secret.id) is None
    assert internal.parent_section(secret.id).text
    result = public.search("SHORT CONDITION")[0]
    assert result.text in public.parent_section(result.id).text
    beta = next(doc for doc in bundle.documents if doc.id == "pol-b")
    assert public.document(beta.version_id) is None
    assert all(node.document_version_id != beta.version_id for node in public.nodes())
    visible_ids = {node.id for node in public.nodes()}
    assert all(
        edge.source_id in visible_ids and edge.target_id in visible_ids
        for edge in public.edges()
    )
    alpha = next(doc for doc in bundle.documents if doc.id == "pol-a")
    graph = traverse(public, alpha.version_id, max_hops=5)
    assert all(node.document_version_id != beta.version_id for node in graph["nodes"])
    assert traverse(public, beta.version_id) == {"nodes": [], "edges": []}


def test_unreviewed_sources_cannot_become_active_just_by_editing_a_flag(tmp_path):
    path = manifest_for(
        tmp_path,
        [
            {
                "id": "example",
                "title": "Example",
                "filenames": ["example.pdf"],
                "review_status": "active",
            }
        ],
    )
    with pytest.raises(ValidationError, match="reviewed_sha256"):
        DocumentManifest.load(path)


def test_failed_rebuild_preserves_the_previous_database(sources, tmp_path):
    bundle = prepare_source_index(*sources)
    path = tmp_path / "index.sqlite"
    write_source_index(bundle, path)
    original = path.read_bytes()
    broken = bundle.model_copy(deep=True)
    broken.chunks[0].text = "not its actual source"
    with pytest.raises(ValueError, match="exact, located"):
        write_source_index(broken, path)
    assert path.read_bytes() == original
    assert PolicyIndex(path).documents()
    assert not list(tmp_path.glob("*.build-*"))


def test_pdf_keeps_physical_pages_and_structured_table_cells(tmp_path):
    directory = tmp_path / "sources"
    directory.mkdir()
    path = directory / "example.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.drawString(50, 740, "I. POLICY")
    for y in [650, 620, 590]:
        pdf.line(50, y, 350, y)
    for x in [50, 200, 350]:
        pdf.line(x, 590, x, 650)
    for x, y, text in [
        (60, 630, "Condition"),
        (210, 630, "Form"),
        (60, 600, "Exception"),
        (210, 600, "FORM-CELL"),
    ]:
        pdf.drawString(x, y, text)
    pdf.showPage()
    pdf.drawString(50, 740, "PAGE TWO REQUIREMENT")
    pdf.save()
    bundle = prepare_source_index(
        directory, manifest_for(tmp_path, [entry(path, "example")])
    )
    document = bundle.documents[0]
    assert {block.location.page for block in document.blocks} == {1, 2}
    table = next(block for block in document.blocks if block.kind == "table")
    assert table.rows[1][1].text == "FORM-CELL"
    assert table.location.bbox and table.location.page == 1
    assert "PAGE TWO REQUIREMENT" in "\n".join(chunk.text for chunk in bundle.chunks)


def test_docx_cannot_fabricate_a_physical_page():
    with pytest.raises(ValidationError, match="physical page"):
        SourceLocation(file="example.docx", format="docx", page=1)


def test_unregistered_source_fails_explicitly_instead_of_silently_skipping(sources):
    new = Document()
    new.add_paragraph("New source")
    new.save(sources[0] / "unregistered.docx")
    with pytest.raises(ValueError, match="Register this source"):
        prepare_source_index(*sources)


def test_source_revisions_are_retained_as_archived_and_not_default_search_results(
    sources, tmp_path
):
    original = prepare_source_index(*sources)
    path = tmp_path / "index.sqlite"
    write_source_index(original, path)
    previous = next(doc for doc in original.documents if doc.id == "pol-a")
    source = sources[0] / "Policy-Alpha.docx"
    doc = Document(source)
    doc.add_paragraph("NEW VERSION MARKER")
    doc.save(source)
    manifest = json.loads(sources[1].read_text())
    manifest["documents"][0]["review_status"] = "review"
    sources[1].write_text(json.dumps(manifest))
    updated = write_source_index(prepare_source_index(*sources), path)
    assert any(
        doc.version_id == previous.version_id and doc.review_status == "archived"
        for doc in updated.documents
    )
    public = PolicyIndex(path)
    assert public.document(previous.version_id) is None
    review = PolicyIndex(path, LOCAL_REVIEW_SCOPE)
    assert review.document(previous.version_id).review_status == "archived"
    assert review.search("NEW VERSION MARKER")
    assert public.search("NEW VERSION MARKER") == []


def test_repeated_pdf_headers_preserve_continuous_section_and_all_page_locations(
    tmp_path,
):
    directory = tmp_path / "sources"
    directory.mkdir()
    path = directory / "example.pdf"
    pdf = canvas.Canvas(str(path))
    for page in range(2):
        pdf.setFont("Helvetica-Bold", 18)
        pdf.drawString(50, 740, "RUNNING DOCUMENT TITLE")
        pdf.setFont("Helvetica", 10)
        pdf.drawString(50, 700, f"Unique body clause on physical page {page + 1}")
        pdf.showPage()
    pdf.save()
    bundle = prepare_source_index(
        directory, manifest_for(tmp_path, [entry(path, "example")])
    )
    titles = [
        section
        for section in bundle.sections
        if section.title == "RUNNING DOCUMENT TITLE"
    ]
    assert len(titles) == 1
    assert "physical page 1" in titles[0].text and "physical page 2" in titles[0].text
    assert {block.location.page for block in titles[0].blocks} == {1, 2}


def test_chroma_records_point_to_exact_parents_and_filter_source_access(sources):
    bundle = prepare_source_index(*sources)
    ids, texts, metadata = chroma_records(bundle)
    assert ids and len(ids) == len(texts) == len(metadata)
    section_ids = {section.id for section in bundle.sections}
    assert all(item["parent_section_id"] in section_ids for item in metadata)
    assert all(
        item["access_classification"] == "public" and item["review_status"] == "active"
        for item in metadata
    )
    assert not any(item["document_id"] == "pol-b" for item in metadata)
    assert all(item["page"] == "" for item in metadata)
    assert len(chroma_records(bundle, LOCAL_REVIEW_SCOPE)[0]) > len(ids)


class FakeTokenizer:
    is_fast = True
    name_or_path = "test-offset-tokenizer"

    def num_special_tokens_to_add(self, pair):
        return 2

    def __call__(self, text, add_special_tokens=True, **kwargs):
        from rag_pipeline.documents.sections import lexical_spans

        spans = lexical_spans(text)
        return {
            "input_ids": list(range(len(spans) + (2 if add_special_tokens else 0))),
            "offset_mapping": spans,
        }


class FakeVector(list):
    def tolist(self):
        return list(self)


class FakeEmbedder:
    tokenizer = FakeTokenizer()
    max_seq_length = 64

    def encode(self, texts, **kwargs):
        return [FakeVector([0.0, 1.0]) for _ in texts]


def test_embedding_chunks_use_actual_offset_mapping_and_model_budget(sources):
    bundle = embedding_chunks(prepare_source_index(*sources), FakeEmbedder())
    assert all(chunk.tokenizer == "test-offset-tokenizer" for chunk in bundle.chunks)
    assert all(chunk.token_count + 2 <= 64 for chunk in bundle.chunks)
    assert "FINAL-TAIL-MARKER" in "\n".join(chunk.text for chunk in bundle.chunks)
    bundle.validate_integrity()


def test_chroma_failure_keeps_active_pointer_and_preview_never_promotes(
    sources, tmp_path, monkeypatch
):
    import sys
    from types import SimpleNamespace

    pointer = tmp_path / "active_policy_collection.json"
    pointer.write_text(json.dumps({"collection": "existing-collection"}))

    class Collection:
        name = "new-collection"
        fail = True
        total = 0

        def add(self, ids, **kwargs):
            if self.fail:
                raise RuntimeError("Simulated partial write")
            self.total += len(ids)

        def count(self):
            return self.total

    collection = Collection()
    client = SimpleNamespace(create_collection=lambda **kwargs: collection)
    monkeypatch.setitem(
        sys.modules,
        "chromadb",
        SimpleNamespace(PersistentClient=lambda **kwargs: client),
    )
    bundle = prepare_source_index(*sources)
    with pytest.raises(RuntimeError, match="partial write"):
        publish_chroma(bundle, FakeEmbedder(), tmp_path)
    assert json.loads(pointer.read_text())["collection"] == "existing-collection"
    collection.fail = False
    publish_chroma(bundle, FakeEmbedder(), tmp_path, preview=True)
    assert json.loads(pointer.read_text())["collection"] == "existing-collection"
    collection.total = 0
    publish_chroma(bundle, FakeEmbedder(), tmp_path)
    assert json.loads(pointer.read_text())["collection"].startswith("uc_policies_")
    snapshot = active_source_index_path(tmp_path)
    assert snapshot.is_file()
    assert PolicyIndex(snapshot).documents()
    assert snapshot != tmp_path / "index.sqlite"


def test_link_and_block_xml_locations_resolve_to_actual_document_nodes(sources):
    from lxml import etree
    import zipfile

    bundle = prepare_source_index(*sources)
    document = next(doc for doc in bundle.documents if doc.id == "pol-a")
    with zipfile.ZipFile(sources[0] / document.file) as package:
        for block in document.blocks:
            root = etree.fromstring(package.read(block.location.part))
            assert root.xpath(
                block.location.element,
                namespaces={
                    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
                },
            )
            for link in block.links:
                assert root.xpath(
                    link.location.element,
                    namespaces={
                        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
                    },
                )


def test_malformed_pdf_annotation_metadata_does_not_lose_source_link_targets():
    import base64
    from types import SimpleNamespace
    from rag_pipeline.documents.extractors import pdf_source_links

    raw = b"\xfe\xff\x00x\x00"
    page = SimpleNamespace(
        height=100,
        page_obj=SimpleNamespace(
            annots=[
                {
                    "Rect": [0, 0, 50, 20],
                    "A": {"URI": b"https://example.invalid"},
                    "Contents": b"\xffmalformed odd annotation metadata",
                },
                {"Rect": [0, 0, 50, 20], "A": {"URI": raw}},
            ]
        ),
        crop=lambda *args, **kwargs: SimpleNamespace(extract_text=lambda: "link label"),
    )
    links = pdf_source_links(
        page, SourceLocation(file="test.pdf", format="pdf", page=1)
    )
    assert links[0].target == "https://example.invalid"
    assert links[1].target.startswith("unresolved-uri:")
    assert base64.b64decode(links[1].raw_target_base64) == raw
